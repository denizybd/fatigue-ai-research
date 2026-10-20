"""Annotate the existing20-life/200-defect data using acquired EAN2025 evidence.

This adds no experiments, fits, physical batch IDs or fabricated covariates.
"""
from pathlib import Path
import collections
import csv
import hashlib
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
LIT = ROOT / 'literature/v5_open/bartosak'
OUT = ROOT / 'data/external_v5/bartosak'
RAW = OUT / 'raw'
workbook = RAW / 'bartosak2024_316L_fatigue_and_defects.xlsx'
wb = load_workbook(workbook, read_only=True, data_only=True)
fatigue = list(wb['Fatigue'].values)
assert fatigue[0] == ('ID', 'PSWT [MPa]', 'Nf [-]', 'Loading path')
doi = '10.1016/j.ijfatigue.2024.108608'
conference = 'http://www.experimentalni-mechanika.cz/cs/konference/konference/2025.html?download=3105%3A21'
life_rows = []
defect_rows = []
for raw_row, values in enumerate(fatigue[1:], 2):
    specimen_id, pswt, nf, loading_path = values
    assert specimen_id == raw_row - 1 and 1 <= specimen_id <= 20
    life_rows.append({
        'author_specimen_label':str(specimen_id),
        'identity_scope':'author workbook specimen label; no verified engraved mark or cross-publication physical identity',
        'source_doi':doi, 'dataset_doi':'10.5281/zenodo.12515773',
        'raw_sheet':'Fatigue', 'raw_row':raw_row,
        'loading_path':loading_path, 'load_control':'strain',
        'regime':'low-cycle fatigue', 'process':'laser powder bed fusion',
        'material':'316L stainless steel', 'geometry':'thin-walled tubular',
        'Nf_cycles':nf, 'event':1, 'runout':0,
        'life_endpoint':'10% equivalent-stress-amplitude drop below tangent to stabilized/linear cyclic evolution',
        'P_SWT_mpa':pswt,
        'P_SWT_definition':'maximum normal stress times normal strain amplitude on maximum-tensile-strain plane',
        'P_SWT_cycle':'mid-life cycle Nf/2',
        'P_SWT_cycle_evidence':'EAN2025 PDF page4 Section5 Eq9 and following sentence',
        'P_SWT_is_prescribed_stress_amplitude':False,
        'P_SWT_free_pretest_input':False,
        'P_SWT_validity_note':'Conditional cyclic-state fatigue model input; not a scientific-error allegation. Prediction-time assumptions must be stated.',
        'temperature_c':'', 'R_strain':'', 'frequency_hz':'',
        'physical_build_id':'', 'powder_lot_id':'',
        'build_crosswalk_verified':False, 'powder_crosswalk_verified':False,
        'gold_build_qualified':False, 'primary_force_axial_HCF_eligible':False,
        'CT_acquired_before_fatigue':True,
        'CT_timing_source_url':conference,
        'CT_timing_source_location':'EAN2025 PDF page2 Section3',
        'source_family':'Bartosak108608_EAN2025_37_42',
        'dataset_license':'CC BY4.0',
    })
    sheet = wb[f'Specimen_{specimen_id}']
    assert sheet.max_row == 11
    for rank, v in enumerate(list(sheet.values)[1:], 1):
        area, sphericity, distance, diameter = v
        defect_rows.append({
            'author_specimen_label':str(specimen_id), 'source_doi':doi,
            'dataset_doi':'10.5281/zenodo.12515773',
            'raw_sheet':sheet.title, 'raw_row':rank+1,
            'defect_row_rank':rank, 'rank_is_author_defect_id':False,
            'projected_area_mm2':area, 'sphericity':sphericity,
            'distance_from_free_surface_mm':distance, 'defect_diameter_mm':diameter,
            'CT_acquired_before_fatigue':True,
            'CT_timing_evidence':'EAN2025 PDF page2 Section3; same20-specimen cohort and explicit108608 reference',
            'projection_plane':'plane experiencing maximum tensile strain',
            'entire_feature_vector_free_pretest_eligibility':'not fully established; raw CT is pretest but critical-plane projection must be consistent with intended loading/state',
            'physical_build_id':'', 'powder_lot_id':'', 'dataset_license':'CC BY4.0',
        })

def write_csv(name, records):
    with (OUT/name).open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(records[0]));w.writeheader();w.writerows(records)

write_csv('specimen20_evidence_annotated.csv',life_rows)
write_csv('defects200_pretest_CT_evidence.csv',defect_rows)
assert len(life_rows)==20 and len(defect_rows)==200
assert collections.Counter(r['loading_path'] for r in life_rows)=={'Proportional':5,'Axial':7,'Torsional':8}
assert not any(r['physical_build_id'] or r['powder_lot_id'] for r in life_rows)
assert all(not r['P_SWT_free_pretest_input'] for r in life_rows)
receipt={
    'status':'PASS', 'existing_life_records_reannotated':20,
    'existing_defect_records_reannotated':200, 'new_numeric_life_records':0,
    'new_source_evidence':'Acquired6-page EAN2025 primary PDF; CT acquisition timing before fatigue and P_SWT evaluation atNf/2 verified',
    'loading_path_counts':dict(collections.Counter(r['loading_path'] for r in life_rows)),
    'gold_build_qualified_records':0, 'primary_force_axial_HCF_eligible_records':0,
    'new_model_fits':0, 'input_workbook_sha256':hashlib.sha256(workbook.read_bytes()).hexdigest(),
    'conference_pdf_sha256':hashlib.sha256((LIT/'conference2025_http.pdf').read_bytes()).hexdigest(),
    'output_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(OUT.glob('*.csv'))},
    'run_command':'/Users/deniz2/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3 literature/v5_open/bartosak/annotate_recovered_evidence.py',
}
(LIT/'normalization_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n')
print(json.dumps(receipt,indent=2))
